"""
build_inventory_workbook.py

Takes the lot crosswalk CSV, calculates landed cost per lot (hammer price +
buyer's premium + tax + allocated shipping), and produces a final Excel
workbook with each lot's photo embedded next to its row.

Usage:
    python build_inventory_workbook.py --csv crosswalk.csv --images ./images \\
        --shipping 19.50 --out inventory.xlsx
"""

import argparse
import os

import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter

from config import (
    BUYERS_PREMIUM_RATE,
    TAX_RATE_ON_PREMIUM,
    IMAGE_COLUMN_LETTER,
    ROW_HEIGHT,
    IMAGE_WIDTH_PX,
    IMAGE_HEIGHT_PX,
)


def compute_costs(df: pd.DataFrame, total_shipping: float) -> pd.DataFrame:
    df["Buyer's Premium"] = BUYERS_PREMIUM_RATE * df["Hammer price"]
    df["Tax"] = TAX_RATE_ON_PREMIUM * df["Buyer's Premium"]

    num_lots = len(df) if len(df) else 1
    df["Shipping (allocated)"] = total_shipping / num_lots

    df["Total Cost"] = (
        df["Hammer price"] + df["Buyer's Premium"] + df["Tax"] + df["Shipping (allocated)"]
    )
    return df


def embed_images(workbook_path: str, df: pd.DataFrame, images_folder: str, sheet_name: str = "Sheet1"):
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb[sheet_name]

    lot_number_col_idx = df.columns.get_loc("Lot number") + 1
    lot_col_letter = get_column_letter(lot_number_col_idx)

    lot_to_row = {}
    for row_idx in range(2, ws.max_row + 1):
        val = ws[f"{lot_col_letter}{row_idx}"].value
        if isinstance(val, (int, float)):
            lot_to_row[int(val)] = row_idx

    for _, record in df.iterrows():
        lot_num = int(record["Lot number"])
        img_name = record.get("Image filename")

        if pd.isna(img_name) or not str(img_name).strip():
            continue
        if lot_num not in lot_to_row:
            print(f"Lot {lot_num} not found in sheet; skipping image.")
            continue

        img_path = os.path.join(images_folder, str(img_name).strip())
        if not os.path.exists(img_path):
            print(f"Image not found for lot {lot_num}: {img_path}")
            continue

        row_idx = lot_to_row[lot_num]
        try:
            ws.row_dimensions[row_idx].height = ROW_HEIGHT
            img = Image(img_path)
            img.width = IMAGE_WIDTH_PX
            img.height = IMAGE_HEIGHT_PX
            ws.add_image(img, f"{IMAGE_COLUMN_LETTER}{row_idx}")
        except Exception as exc:
            print(f"Error inserting image for lot {lot_num}: {exc}")

    wb.save(workbook_path)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Build a financially-modeled inventory workbook with embedded photos.")
    parser.add_argument("--csv", required=True, help="Path to the lot crosswalk CSV")
    parser.add_argument("--images", required=True, help="Folder containing lot images")
    parser.add_argument("--shipping", type=float, required=True, help="Total shipping cost to allocate across all lots")
    parser.add_argument("--out", required=True, help="Output .xlsx path")
    args = parser.parse_args()

    lots_df = pd.read_csv(args.csv)
    lots_df = compute_costs(lots_df, args.shipping)
    lots_df.to_excel(args.out, index=False)
    print(f"Saved cost-modeled workbook to {args.out}")

    embed_images(args.out, lots_df, args.images)
    print(f"Embedded lot images into {args.out}")
